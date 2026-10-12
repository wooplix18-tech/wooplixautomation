"""Export the saved artifact without regenerating or contacting a model."""

import io
import json
import tempfile
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document

from .models import Artifact


def artifact_docx(artifact: Artifact) -> bytes:
    if artifact.content.get("integration", {}).get("engine") == "proposalagent-langgraph":
        from .proposal_agent import business_requirements_agent, wooplix_agent

        renderer = wooplix_agent.build_docx if artifact.kind == "proposal" else business_requirements_agent.build_docx
        with tempfile.TemporaryDirectory(prefix="wooplix-export-") as folder:
            output = Path(folder) / "document.docx"
            renderer(artifact.content, str(output))
            return output.read_bytes()
    document = Document()
    document.add_heading(artifact.title, 0)
    document.add_paragraph("Wooplix Technologies Private Limited")
    document.add_paragraph(f"Draft · version {artifact.version}")
    if artifact.kind == "documentation":
        content = artifact.content
        for key, value in content.items():
            if key == "document_status":
                document.add_paragraph(str(value))
                continue
            document.add_heading(key.replace("_", " ").title(), 1)
            _write_docx_value(document, value)
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()
    for key, value in artifact.content.items():
        if key in {"planning_target", "traditional_hours", "ai_assisted_hours", "human_involvement_hours"} or key.endswith("_hours_target"):
            continue
        document.add_heading(key.replace("_", " ").title(), 1)
        document.add_paragraph(value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, indent=2))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def artifact_pdf(artifact: Artifact) -> bytes:
    """Render a saved artifact through Dompdf, locally or on Vercel."""
    import json
    import os
    import subprocess
    import urllib.error
    import urllib.request

    payload = json.dumps(artifact.model_dump(mode="json"), ensure_ascii=False).encode("utf-8")
    # Python functions on Vercel cannot assume a PHP binary is installed. Use
    # the companion PHP function, which loads the same Composer Dompdf package.
    deployment = os.getenv("VERCEL_URL", "").strip()
    if os.getenv("VERCEL") == "1" and deployment:
        request = urllib.request.Request(
            f"https://{deployment}/api/render-pdf.php",
            data=payload,
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        try:
            with urllib.request.urlopen(request, timeout=90) as response:
                if response.status != 200 or response.headers.get_content_type() != "application/pdf":
                    raise RuntimeError("Dompdf could not render this document.")
                return response.read()
        except (urllib.error.URLError, TimeoutError) as exc:
            raise RuntimeError("The Dompdf PDF service could not be reached.") from exc

    php = __import__("shutil").which("php")
    script = Path(__file__).resolve().parent.parent / "api" / "render-pdf.php"
    autoload = script.parent.parent / "vendor" / "autoload.php"
    if not php or not autoload.is_file():
        raise RuntimeError("Dompdf is unavailable. Install PHP dependencies with composer install.")
    result = subprocess.run([php, str(script)], input=payload, capture_output=True, timeout=90)
    if result.returncode or not result.stdout.startswith(b"%PDF-"):
        detail = result.stderr.decode("utf-8", errors="replace")[:500]
        raise RuntimeError(f"Dompdf could not render this document. {detail}")
    return result.stdout

def _write_docx_value(document: Document, value, depth: int = 0) -> None:
    """Render nested documentation as readable Word sections and lists."""
    if isinstance(value, dict):
        for key, child in value.items():
            label = key.replace("_", " ").capitalize()
            if isinstance(child, (dict, list)):
                if depth < 2:
                    document.add_heading(label, min(depth + 2, 3))
                else:
                    paragraph = document.add_paragraph()
                    paragraph.add_run(label + ":").bold = True
                _write_docx_value(document, child, depth + 1)
            else:
                paragraph = document.add_paragraph()
                paragraph.add_run(label + ": ").bold = True
                paragraph.add_run(str(child) if child not in (None, "") else "To be confirmed")
    elif isinstance(value, list):
        for item in value:
            if isinstance(item, dict):
                title = item.get("requirement_id") or item.get("audience") or item.get("scenario")
                if isinstance(title, list):
                    title = ", ".join(str(part) for part in title)
                if title:
                    document.add_heading(str(title), min(depth + 2, 3))
                _write_docx_value(document, item, depth + 1)
            elif isinstance(item, (list, tuple)):
                _write_docx_value(document, list(item), depth + 1)
            else:
                document.add_paragraph(str(item), style="List Bullet")
    elif value not in (None, ""):
        document.add_paragraph(str(value))
