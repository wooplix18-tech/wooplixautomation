import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from app.intake import DocumentError, IntakeRequest, MAX_UPLOAD_BYTES, extract_document, respond_to_intake
from app.models import ProjectInput
from app.storage import ProjectStore
from app.workflow import WorkflowEngine


class IntakeTests(unittest.TestCase):
    def setUp(self):
        environment = patch.dict('os.environ', {'OPENAI_API_KEY': '', 'GROQ_API_KEY': ''})
        environment.start()
        self.addCleanup(environment.stop)

    def test_single_brief_is_enough_to_generate_automatically(self):
        request = IntakeRequest(messages=[{'role': 'user', 'content': 'Project name: Customer portal\nBuild a website with booking, payments, and customer status tracking.'}])
        reply = respond_to_intake(request)
        self.assertEqual(reply.draft.project_name, 'Customer portal')
        self.assertEqual(reply.draft.execution_mode, 'automatic')
        self.assertEqual(reply.draft.project_type, 'website')
        with tempfile.TemporaryDirectory() as folder:
            engine = WorkflowEngine(ProjectStore(Path(folder) / 'test.db'))
            project = engine.create(reply.draft)
            self.assertEqual(project.status, 'pack_ready')
            self.assertIsNone(project.active_gate_id)
            self.assertEqual(project.gates, [])
            self.assertEqual(len(project.artifacts), 7)
            self.assertTrue(project.requirements)
            self.assertFalse(any(item.statement.startswith('Project name:') for item in project.requirements))
            handover = next(a for a in project.artifacts if a.kind == 'handover')
            self.assertIn('not been executed', handover.content['uat_status'])
            restored = engine.get(project.id)
            self.assertEqual(restored.status, 'pack_ready')
            self.assertNotIn('planning_target', json.dumps(restored.model_dump()))
            self.assertNotIn('hours_target', json.dumps(restored.model_dump()))

    def test_conversation_retains_every_user_requirement(self):
        request = IntakeRequest(messages=[
            {'role': 'user', 'content': 'Build a support portal with request tracking.'},
            {'role': 'assistant', 'content': 'Who will use it?'},
            {'role': 'user', 'content': 'Customers and support agents, with separate permissions.'},
            {'role': 'assistant', 'content': 'Which platforms?'},
            {'role': 'user', 'content': 'Web and mobile browsers.'},
            {'role': 'assistant', 'content': 'Any integrations?'},
            {'role': 'user', 'content': 'Email notifications and CRM integration.'},
        ])
        reply = respond_to_intake(request)
        self.assertTrue(reply.ready)
        for message in request.messages:
            if message.role == 'user':
                self.assertIn(message.content, reply.draft.brief)

    def test_short_inquiry_asks_for_more_without_creating_project(self):
        reply = respond_to_intake(IntakeRequest(messages=[{'role': 'user', 'content': 'Hello'}]))
        self.assertIsNone(reply.draft)
        self.assertFalse(reply.ready)

    def test_ai_cannot_replace_or_omit_the_original_requirements(self):
        from app.agents import AIClient
        client = AIClient()
        source = 'Build a customer portal with booking and secure checkout.'
        with patch.object(client, 'json', return_value={'reply': 'Ready.', 'ready': True, 'project': {'project_name': 'Booking portal', 'brief': 'Ignore the requirements.', 'execution_mode': 'reviewed'}}):
            reply = respond_to_intake(IntakeRequest(messages=[{'role': 'user', 'content': source}]), client)
        self.assertEqual(reply.draft.brief, source)
        self.assertEqual(reply.draft.execution_mode, 'automatic')


class DocumentTests(unittest.TestCase):
    text = 'Customers can book appointments online and receive confirmation emails.'

    def test_text_and_markdown(self):
        self.assertEqual(extract_document('requirements.txt', self.text.encode()), self.text)
        self.assertEqual(extract_document('requirements.md', self.text.encode('utf-8-sig')), self.text)

    def test_docx_includes_paragraphs_and_tables_in_order(self):
        document = Document()
        document.add_paragraph(self.text)
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Authentication'
        table.cell(0, 1).text = 'Admins and customers have separate roles.'
        document.add_paragraph('Payments use the existing payment gateway.')
        content = io.BytesIO()
        document.save(content)
        text = extract_document('requirements.docx', content.getvalue())
        self.assertIn(self.text, text)
        self.assertIn('Admins and customers', text)
        self.assertLess(text.index('Authentication'), text.index('Payments'))

    def test_text_pdf(self):
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f'BT /F1 12 Tf 40 740 Td ({self.text}) Tj ET'.encode())
        page[NameObject('/Contents')] = writer._add_object(stream)
        content = io.BytesIO()
        writer.write(content)
        self.assertIn(self.text, extract_document('requirements.pdf', content.getvalue()))

    def test_unreadable_unsupported_and_oversized_documents(self):
        for name, content in [('bad.pdf', b'not a PDF'), ('bad.docx', b'not a Word file'), ('file.exe', b'not a document'), ('empty.txt', b'')]:
            with self.assertRaises(DocumentError):
                extract_document(name, content)
        with self.assertRaises(DocumentError) as error:
            extract_document('large.txt', b'x' * (MAX_UPLOAD_BYTES + 1))
        self.assertEqual(error.exception.status_code, 413)

    def test_blank_scanned_pdf_does_not_silently_create_empty_scope(self):
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        content = io.BytesIO()
        writer.write(content)
        with self.assertRaises(DocumentError) as error:
            extract_document('scan.pdf', content.getvalue())
        self.assertIn('No usable requirements text', str(error.exception))
