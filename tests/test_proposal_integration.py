import io
import unittest
from unittest.mock import patch, MagicMock

from docx import Document

from app.exports import artifact_docx
from app.models import Project, ProjectInput, RequirementItem
from app.proposal_integration import IntegratedDocumentAgent, proposal_engine


class ProposalIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.project = Project(input=ProjectInput(project_name="CRM project", brief="Configure Zoho CRM with lead management."),
                               requirements=[RequirementItem(id="REQ-1", category="Zoho CRM", statement="Lead management")])

    def test_explicit_live_engine_requires_provider(self):
        with patch.dict('os.environ', {'PROPOSAL_ENGINE':'langgraph', 'GROQ_API_KEY':''}):
            with self.assertRaises(ValueError):
                proposal_engine()

    def test_default_uses_offline_template_without_key(self):
        with patch.dict('os.environ', {'PROPOSAL_ENGINE':'auto', 'GROQ_API_KEY':''}):
            self.assertEqual(proposal_engine(), 'template')

    def test_quality_failure_cannot_be_persisted_as_success(self):
        graph = MagicMock()
        graph.invoke.return_value = {'status':'rejected', 'proposal':{}, 'quality_report':{'passed':False}}
        with patch('app.proposal_agent.multi_agent_graph.build_graph', return_value=graph):
            with self.assertRaises(ValueError):
                IntegratedDocumentAgent('proposal').run(self.project)

    def test_verified_saved_proposal_exports_with_original_renderer(self):
        proposal = {'client':{'company_name':'Client', 'project_name':'CRM', 'contact':''},
                    'project_introduction':'Implement Zoho CRM.',
                    'scope':[{'product':'Zoho CRM', 'areas':[{'area':'Leads', 'tasks':['Lead management']}]}],
                    'prerequisites':[], 'deliverables':['Configured CRM'], 'open_points':[], 'status':'DRAFT'}
        graph = MagicMock()
        graph.invoke.return_value = {'status':'draft_ready', 'proposal':proposal, 'quality_report':{'passed':True}, 'trace':['supervisor', 'critic passed']}
        with patch('app.proposal_agent.multi_agent_graph.build_graph', return_value=graph):
            artifact = IntegratedDocumentAgent('proposal').run(self.project)
        self.assertIn('quality_critic', artifact.generated_by)
        self.assertEqual(artifact.content['shared_requirement_links'][0]['id'], 'REQ-1')
        data = artifact_docx(artifact)
        document = Document(io.BytesIO(data))
        text = '\n'.join(p.text for p in document.paragraphs) + '\n'.join(cell.text for t in document.tables for r in t.rows for cell in r.cells)
        self.assertIn('Lead management', text)

    def test_source_ids_remain_stable_when_lines_reordered(self):
        from app.agents import _requirement_id
        self.assertEqual(_requirement_id('Customers can book appointments.',1), _requirement_id('Customers can book appointments.',8))
